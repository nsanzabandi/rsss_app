"""
core/exports.py — Downloadable child lists, limited to the user's catchment.

Every list is scoped twice, server-side:
  1. by the signed-in user's role (ministry → all; district / hospital /
     health centre → only their own area — same rule as data.filter_by_user)
  2. by the dashboard's current province / district / hospital filters
     (and date range, for the period-based lists).

Files contain children's names and parents' phone numbers (needed for
follow-up), so every download is appended to logs/downloads.log and the
workbook opens with an "About" sheet stating scope and confidentiality.
"""
from __future__ import annotations

import io
import json
import threading
from datetime import datetime, timedelta
from pathlib import Path

import pandas as pd

BASE_DIR  = Path(__file__).resolve().parent.parent
_LOG_PATH = BASE_DIR / "logs" / "downloads.log"
_log_lock = threading.Lock()

MAX_ROWS = 300_000          # beyond this the file is too big to send to a browser
MISSED_DAYS, MISSED_LOOKBACK = 10, 180

LISTS = {
    "stunted":  "Stunted children",
    "at_risk":  "At-risk children",
    "missed":   "Missed appointments",
    "assessed": "All assessed children",
}

_DIST = "COALESCE(h_district, residence_district, district_source)"

# DB columns pulled for every list (latest relevant visit per child)
_DETAIL_SQL = f"""
SELECT entity_id                AS tracked_entity_instance,
       event_id, child_id, child_name, gender, date_of_birth,
       last_immunization_date   AS visit_date,
       age_visit_months         AS age_months,
       height_visit_cm          AS height_cm,
       weight_visit_kg          AS weight_kg,
       muac_cm, wasting_status, immunization_schedule, next_visit_date,
       residence_province       AS province,
       {_DIST}                  AS district,
       h_district_hospital      AS district_hospital,
       health_facility, h_sector AS sector, village,
       mother_names, mother_phone, father_names, father_phone
FROM immunization_vaccination
"""

# (column, header) — order of columns in the downloaded sheet
_COLUMNS = [
    ("child_id", "Child ID"), ("child_name", "Child name"), ("gender", "Sex"),
    ("date_of_birth", "Date of birth"), ("age_months", "Age (months)"),
    ("status", "Status"), ("risk_level", "Risk level"), ("risk_reason", "Why at risk"),
    ("days_overdue", "Days overdue"), ("next_visit_date", "Next visit due"),
    ("visit_date", "Latest visit"), ("height_cm", "Height (cm)"), ("weight_kg", "Weight (kg)"),
    ("muac_cm", "MUAC (cm)"), ("weight_velocity", "Weight change (kg/month)"),
    ("immunization_schedule", "Schedule"),
    ("province", "Province"), ("district", "District"), ("district_hospital", "Hospital"),
    ("health_facility", "Health facility"), ("sector", "Sector"), ("village", "Village"),
    ("mother_names", "Mother"), ("mother_phone", "Mother phone"),
    ("father_names", "Father"), ("father_phone", "Father phone"),
]


class ExportTooLarge(Exception):
    pass


# ── Scope ──────────────────────────────────────────────────────────────────────

def scope_label(user: dict, province=None, district=None, hospital=None) -> str:
    role = user.get("role")
    own = {"district": user.get("district"), "hospital": user.get("hospital"),
           "health_center": user.get("health_center")}.get(role)
    parts = [p for p in (own, province, district, hospital) if p]
    return " / ".join(dict.fromkeys(parts)) or "Rwanda (national)"


def _sql_scope(user: dict, province=None, district=None, hospital=None) -> tuple[str, list]:
    """WHERE clauses for the user's role + dashboard filters (SQL lists)."""
    where, params = [], []
    role = user.get("role")
    if role == "district":
        where.append(f"{_DIST} = %s"); params.append(user.get("district"))
    elif role == "hospital":
        where.append("h_district_hospital = %s"); params.append(user.get("hospital"))
    elif role == "health_center":
        where.append("health_facility = %s"); params.append(user.get("health_center"))
    elif role != "ministry":
        where.append("FALSE")                       # unknown role → nothing
    if province:
        where.append("residence_province = %s"); params.append(province)
    if district:
        where.append(f"{_DIST} = %s"); params.append(district)
    if hospital:
        where.append("h_district_hospital = %s"); params.append(hospital)
    return " AND ".join(where), params


def _scoped_children(user, province, district, hospital, start, end) -> pd.DataFrame:
    """Per-child computed status (same cohort as the dashboard KPI cards)."""
    import data
    child, visits = data.get_child_df(), data.get_visits_df()
    if child is None:
        raise RuntimeError("Dashboard data is still loading — try again in a minute.")
    return data.scoped_child_df(child, visits, user, province, district, hospital, start, end)


def _details(ids: list, start=None, end=None) -> pd.DataFrame:
    """Latest visit (within the date window) per child, with contact details."""
    from config.db_local import get_local_conn
    if not ids:
        return pd.DataFrame()
    sql, params = _DETAIL_SQL + " WHERE entity_id = ANY(%s)", [ids]
    if start:
        sql += " AND last_immunization_date >= %s"; params.append(start)
    if end:
        sql += " AND last_immunization_date <= %s"; params.append(end)
    conn = get_local_conn(); cur = conn.cursor()
    cur.execute(sql, params)
    cols = [d[0] for d in cur.description]
    df = pd.DataFrame(cur.fetchall(), columns=cols)
    cur.close(); conn.close()
    if df.empty:
        return df
    return (df.sort_values(["visit_date", "event_id"])
              .drop_duplicates("tracked_entity_instance", keep="last"))


def _check_size(n: int) -> None:
    if n > MAX_ROWS:
        raise ExportTooLarge(f"{n:,} children — too many for one file. "
                             f"Pick a province or district first.")


# ── Lists ──────────────────────────────────────────────────────────────────────

def _stunted(user, province, district, hospital, start, end, **_) -> pd.DataFrame:
    c = _scoped_children(user, province, district, hospital, start, end)
    c = c[c["is_stunted"].eq(True)]
    _check_size(len(c))
    d = _details(c["tracked_entity_instance"].tolist(), start, end)
    severe = set(c.loc[c["is_severe"].eq(True), "tracked_entity_instance"])
    if not d.empty:
        d["status"] = d["tracked_entity_instance"].map(
            lambda t: "Severe stunting" if t in severe else "Moderate stunting")
        d = d.sort_values(["status", "district", "district_hospital", "health_facility"],
                          ascending=[False, True, True, True])
    return d


def _assessed(user, province, district, hospital, start, end, **_) -> pd.DataFrame:
    c = _scoped_children(user, province, district, hospital, start, end)
    c = c[c["is_stunted"].notna()]
    _check_size(len(c))
    d = _details(c["tracked_entity_instance"].tolist(), start, end)
    if not d.empty:
        st = c.set_index("tracked_entity_instance")
        d["status"] = d["tracked_entity_instance"].map(
            lambda t: ("Severe stunting" if st.at[t, "is_severe"] else
                       "Moderate stunting" if st.at[t, "is_stunted"] else "Normal")
            if t in st.index else "")
        d = d.sort_values(["district", "district_hospital", "health_facility", "child_name"])
    return d


def _at_risk(user, province, district, hospital, start, end, stunted_only=False) -> pd.DataFrame:
    """HIGH + MEDIUM growth risk at each child's latest weighed visit in the
    period — the same definition as the at-risk dashboard (data.scoped_risk_df).
    stunted_only → only children who were also stunted at that visit."""
    import data
    from core.risk_classifier import REASONS
    risk = data.get_risk_df()
    if risk is None:
        raise RuntimeError("Dashboard data is still loading — try again in a minute.")
    r = data.scoped_risk_df(risk, user, province, district, hospital, start, end)
    if r.empty:
        return pd.DataFrame()
    r = r[r["risk_level"].isin(["HIGH", "MEDIUM"])]
    if stunted_only and "is_stunted" in r.columns:
        r = r[r["is_stunted"].eq(True)]
    _check_size(len(r))
    d = _details(r["tracked_entity_instance"].tolist())        # their latest visit + contacts
    if d.empty:
        return d
    extra = r[["tracked_entity_instance", "risk_level", "risk_flags", "weight_velocity"]].copy()
    extra["risk_level"] = extra["risk_level"].astype(object)
    d = d.merge(extra, on="tracked_entity_instance", how="left")
    d["risk_reason"] = d["risk_flags"].fillna(0).astype(int).map(
        lambda f: ", ".join(label for bit, label in REASONS.items() if f & bit) or "—")
    d["weight_velocity"] = pd.to_numeric(d["weight_velocity"], errors="coerce").round(2)
    d["_o"] = d["risk_level"].map({"HIGH": 0, "MEDIUM": 1})
    return d.sort_values(["_o", "weight_velocity"], na_position="last").drop(columns="_o")


def _missed(user, province, district, hospital, start, end, **_) -> pd.DataFrame:
    """Children whose next visit was due 10–180 days ago, by their latest
    visit. Current state — the dashboard date range does not apply."""
    from config.db_local import get_local_conn
    where, params = _sql_scope(user, province, district, hospital)
    today = datetime.now().date()
    sql = (f"SELECT * FROM (SELECT DISTINCT ON (tracked_entity_instance) * FROM ({_DETAIL_SQL}"
           + (f" WHERE {where}" if where else "")
           + ") v ORDER BY tracked_entity_instance, visit_date DESC, event_id DESC) latest "
           "WHERE next_visit_date BETWEEN %s AND %s")
    params += [today - timedelta(days=MISSED_LOOKBACK), today - timedelta(days=MISSED_DAYS)]
    conn = get_local_conn(); cur = conn.cursor()
    cur.execute(sql, params)
    cols = [d[0] for d in cur.description]
    df = pd.DataFrame(cur.fetchall(), columns=cols)
    cur.close(); conn.close()
    _check_size(len(df))
    if not df.empty:
        df["days_overdue"] = (pd.Timestamp(today) - pd.to_datetime(df["next_visit_date"])).dt.days
        df = df.sort_values("days_overdue", ascending=False)
    return df


_BUILDERS = {"stunted": _stunted, "at_risk": _at_risk, "missed": _missed, "assessed": _assessed}


# ── File + audit ───────────────────────────────────────────────────────────────

def _to_xlsx(df: pd.DataFrame, title: str, about: list[tuple[str, str]]) -> bytes:
    """Fast write-only workbook: an 'About' sheet, then the list."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    wb = Workbook(write_only=True)
    info = wb.create_sheet("About")
    info.column_dimensions["A"].width, info.column_dimensions["B"].width = 22, 90
    bold = Font(bold=True)
    for k, v in about:
        cell = WriteOnlyCell(info, value=k); cell.font = bold
        info.append([cell, v])

    ws = wb.create_sheet(title[:31])
    cols = [(c, h) for c, h in _COLUMNS if c in df.columns]
    hfont, hfill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0B2C4A")
    widths = {"child_name": 26, "health_facility": 24, "district_hospital": 24, "mother_names": 24,
              "father_names": 24, "risk_reason": 30, "village": 18}
    for i, (c, _) in enumerate(cols):
        ws.column_dimensions[get_column_letter(i + 1)].width = widths.get(c, 14)
    ws.freeze_panes = "A2"
    head = []
    for _, h in cols:
        cell = WriteOnlyCell(ws, value=h); cell.font, cell.fill = hfont, hfill
        head.append(cell)
    ws.append(head)
    out = df[[c for c, _ in cols]].copy()
    for c in out.columns:
        if pd.api.types.is_datetime64_any_dtype(out[c]):
            out[c] = out[c].dt.date
    out = out.astype(object).where(out.notna(), None)
    def clean(v):
        if isinstance(v, str):
            return ILLEGAL_CHARACTERS_RE.sub("", v)    # stray control chars from eTracker text
        return float(v) if hasattr(v, "is_finite") else v   # Decimal → float
    for row in out.itertuples(index=False, name=None):
        ws.append([clean(v) for v in row])

    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def _audit(user: dict, kind: str, scope: str, filters: dict, rows: int, ok: bool, note: str = "") -> None:
    _LOG_PATH.parent.mkdir(exist_ok=True)
    line = json.dumps({"ts": datetime.now().isoformat(timespec="seconds"),
                       "user": user.get("username"), "role": user.get("role"),
                       "list": kind, "scope": scope, "filters": filters,
                       "rows": rows, "ok": ok, "note": note}, ensure_ascii=False)
    with _log_lock, open(_LOG_PATH, "a") as f:
        f.write(line + "\n")


def build_download(kind: str, user: dict, province=None, district=None, hospital=None,
                   start=None, end=None, stunted_only: bool = False) -> tuple[bytes, str, int]:
    """Return (xlsx bytes, filename, row count). Raises ExportTooLarge /
    RuntimeError with a user-facing message."""
    if kind not in _BUILDERS:
        raise RuntimeError("Unknown list.")
    if not user or user.get("readonly"):
        raise RuntimeError("Downloads aren't available for view-only links.")
    scope = scope_label(user, province, district, hospital)
    filters = {"province": province, "district": district, "hospital": hospital,
               "start": start, "end": end, "stunted_only": stunted_only}
    try:
        df = _BUILDERS[kind](user, province, district, hospital, start, end,
                             stunted_only=stunted_only)
    except Exception as exc:
        _audit(user, kind, scope, filters, 0, False, str(exc)[:200])
        raise
    title = LISTS[kind]
    period = ("current (latest visit)" if kind == "missed"
              else f"{start or 'start'} to {end or 'today'}")
    about = [
        ("List", title + (" — stunted children only" if stunted_only and kind == "at_risk" else "")),
        ("Area", scope),
        ("Period", period if kind != "missed" else
         f"next visit due {MISSED_LOOKBACK}–{MISSED_DAYS} days ago"),
        ("Children", f"{len(df):,}"),
        ("Generated", datetime.now().strftime("%d %b %Y %H:%M")),
        ("Generated by", f"{user.get('full_name') or user.get('username')} ({user.get('role')})"),
        ("Source", "RSSS — NHIC, Ministry of Health Rwanda (eTracker data; WHO height-for-age)"),
        ("CONFIDENTIAL", "Contains children's and parents' personal data. Use only for "
                         "follow-up care; do not forward or share outside your team."),
    ]
    data_bytes = _to_xlsx(df, title, about)
    safe_scope = "".join(ch if ch.isalnum() else "_" for ch in scope).strip("_")[:40]
    fname = f"RSSS_{kind}_{safe_scope}_{datetime.now():%Y%m%d}.xlsx"
    _audit(user, kind, scope, filters, len(df), True)
    return data_bytes, fname, len(df)
