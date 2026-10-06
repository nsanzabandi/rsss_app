"""
core/report_generator.py — PDF and Excel report generation for hospitals/facilities.

Uses fpdf2 for PDF, openpyxl for Excel.
"""
from __future__ import annotations

import os
from datetime import datetime
from typing import TYPE_CHECKING

import openpyxl
from openpyxl.styles import (
    Alignment, Border, Font, PatternFill, Side
)
from openpyxl.utils import get_column_letter
from fpdf import FPDF
from fpdf.enums import XPos, YPos

if TYPE_CHECKING:
    from utils.logger import ReportLogger

# ── Unicode → Latin-1 sanitiser (fpdf2 uses Helvetica which is Latin-1 only) ──

_UNICODE_MAP = str.maketrans({
    "—": " - ",   # em dash  —
    "–": " - ",   # en dash  –
    "‒": "-",     # figure dash
    "‐": "-",     # hyphen
    "‘": "'",     # left single quote  '
    "’": "'",     # right single quote '
    "“": '"',     # left double quote  "
    "”": '"',     # right double quote "
    "…": "...",   # ellipsis  …
    "°": " deg",  # degree sign
    "±": "+/-",   # plus-minus
    "×": "x",     # multiplication sign
    "÷": "/",     # division sign
    "\xa0":   " ",     # non-breaking space
})


def _safe(text) -> str:
    """Strip characters outside Latin-1 so Helvetica never errors."""
    s = str(text) if text is not None else ""
    s = s.translate(_UNICODE_MAP)
    # Remove any remaining non-Latin-1 characters
    return s.encode("latin-1", errors="replace").decode("latin-1")


# Brand colours
_RED_HEX   = "C0392B"
_BLUE_HEX  = "2980B9"
_DARK_HEX  = "1B2A4A"
_GREY_HEX  = "7F8C8D"
_LIGHT_HEX = "F0F2F5"

MONTH_NAMES = {
    1: "January", 2: "February", 3: "March", 4: "April",
    5: "May", 6: "June", 7: "July", 8: "August",
    9: "September", 10: "October", 11: "November", 12: "December",
}

# Excel columns in the case-list output (matches the standard RSSS case list)
CASE_LIST_COLS = [
    ("child_id",                "Child ID"),
    ("child_name",              "Child Name"),
    ("date_of_birth",           "Date of Birth"),
    ("gender",                  "Gender"),
    ("age_in_months",           "Age (Months)"),
    ("mother_names",            "Mother Names"),
    ("mother_phone",            "Mother Phone"),
    ("mother_dob",              "Mother Date of Birth"),
    ("father_names",            "Father Names"),
    ("father_phone",            "Father Phone"),
    ("province",                "Province"),
    ("district",                "District"),
    ("district_hospital",       "District Hospital"),
    ("sector",                  "Sector"),
    ("health_facility",         "Health Facility"),
    ("residence_province",      "Residence Province"),
    ("residence_district",      "Residence District"),
    ("residence_sector",        "Residence Sector"),
    ("village",                 "Village"),
    ("immunization_date",       "Last Immunization Date"),
    ("immunization_schedule",   "Immunization Schedule"),
    ("weight_at_visit_kg",      "Weight (kg)"),
    ("height_at_visit_cm",      "Height (cm)"),
]


# ── PDF ────────────────────────────────────────────────────────────────────────

class _PDF(FPDF):
    def __init__(self, title: str, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._title = title
        self.set_auto_page_break(auto=True, margin=15)
        self.set_margins(15, 15, 15)

    # fpdf2 2.8 removed the `ln=` parameter. Keep supporting `ln=True/False` by
    # translating it to the modern new_x/new_y API, so all call sites keep working.
    def cell(self, *args, ln=None, **kwargs):
        if ln is not None:
            if ln:
                kwargs.setdefault("new_x", XPos.LMARGIN)
                kwargs.setdefault("new_y", YPos.NEXT)
            else:
                kwargs.setdefault("new_x", XPos.RIGHT)
                kwargs.setdefault("new_y", YPos.TOP)
        return super().cell(*args, **kwargs)

    def header(self):
        # Clean report — no letterhead banner (matches the standard RSSS layout).
        pass

    def footer(self):
        self.set_y(-12)
        self.set_font("Helvetica", "", 8)
        self.set_text_color(150, 150, 150)
        self.cell(0, 5, f"Page {self.page_no()}/{{nb}}   |   CONFIDENTIAL", align="C")
        self.set_text_color(0, 0, 0)


class HospitalReportGenerator:
    def __init__(self, log: "ReportLogger"):
        self.log = log

    def create_pdf_report(
        self,
        name: str,
        district: str,
        stats: dict,
        charts: dict,
        year: int,
        month: int,
        output_path: str,
        total_vaccinated: int = 0,
        report_level: str = "hospital",
        total_assessed: int = 0,
    ) -> bool:
        """
        Create a PDF stunting report (standard RSSS layout) for a hospital or
        health facility. Returns True on success.
        """
        try:
            from config.app_config import SUPPORT_PHONE, SUPPORT_CONTACT

            is_hosp     = report_level == "hospital"
            title_txt   = ("DISTRICT HOSPITAL STUNTING REPORT" if is_hosp
                           else "HEALTH FACILITY STUNTING REPORT")
            ent_label   = "Hospital" if is_hosp else "Facility"
            period      = f"{MONTH_NAMES.get(month, str(month))} {year}"
            report_date = datetime.now().strftime("%B %d, %Y")
            total       = stats.get("total_stunted_children", 0)
            avg_age     = stats.get("avg_age_months")
            sched       = stats.get("schedule_breakdown", {})
            vacc        = total_vaccinated or 0
            # Rate is over ASSESSED children (usable height/DOB) — same as the
            # dashboard and the trend chart, so all three show the same %.
            denom       = total_assessed or vacc
            pct_vacc    = round(total / denom * 100, 1) if denom else 0
            safe_name   = _safe(name)
            safe_dist   = _safe(district)
            RED         = (192, 57, 43)

            pdf = _PDF(title=f"{ent_label} Report - {safe_name}")
            pdf.alias_nb_pages()
            pdf.add_page()

            # ── Title ─────────────────────────────────────────────────────────
            pdf.set_text_color(*RED)
            pdf.set_font("Helvetica", "B", 18)
            pdf.cell(0, 12, _safe(title_txt), ln=True, align="C")
            pdf.ln(1)
            pdf.set_text_color(0, 0, 0)
            pdf.set_font("Helvetica", "B", 13)
            pdf.cell(0, 8, _safe(f"{ent_label}: {safe_name}"), ln=True, align="C")
            pdf.cell(0, 8, _safe(f"District: {safe_dist}"), ln=True, align="C")
            pdf.set_text_color(*RED)
            pdf.set_font("Helvetica", "", 10)
            pdf.cell(0, 6, _safe(f"Report Period: {period}"), ln=True, align="C")
            pdf.cell(0, 6, _safe(f"Report Date: {report_date}"), ln=True, align="C")
            pdf.set_text_color(0, 0, 0)
            pdf.ln(4)

            # ── Immediate action box ──────────────────────────────────────────
            bx, bw = 15, 180
            y0 = pdf.get_y()
            pdf.set_draw_color(*RED); pdf.set_line_width(0.6)
            pdf.set_fill_color(253, 237, 236)
            pdf.rect(bx, y0, bw, 26, "DF")
            pdf.set_xy(bx, y0 + 3)
            pdf.set_text_color(*RED); pdf.set_font("Helvetica", "B", 12)
            pdf.cell(bw, 7, "IMMEDIATE ACTION REQUIRED", ln=True, align="C")
            pdf.set_x(bx); pdf.set_font("Helvetica", "", 10)
            pdf.cell(bw, 6, _safe(f"Identified {total:,} stunted out of {denom:,} "
                                  f"children assessed ({pct_vacc}%)"), ln=True, align="C")
            pdf.set_x(bx)
            pdf.cell(bw, 6, "Provide nutrition counseling and coordinate follow-up",
                     ln=True, align="C")
            pdf.set_text_color(0, 0, 0); pdf.set_line_width(0.2)
            pdf.set_y(y0 + 31)

            # ── Key statistics ────────────────────────────────────────────────
            pdf.set_font("Helvetica", "B", 13)
            pdf.cell(0, 8, "Key Statistics", ln=True)
            pdf.ln(1)
            pdf.set_font("Helvetica", "", 11)
            pdf.cell(0, 6, _safe(f"Total Stunted Children: {total:,} out of "
                                 f"{denom:,} assessed ({pct_vacc}%)"), ln=True)
            if vacc:
                pdf.cell(0, 6, _safe(f"Children Vaccinated This Period: {vacc:,}"), ln=True)
            if avg_age:
                pdf.cell(0, 6, _safe(f"Average Age: {avg_age} months"), ln=True)
            if sched:
                top_s, top_n = max(sched.items(), key=lambda x: x[1])
                top_p = round(top_n / total * 100, 1) if total else 0
                pdf.cell(0, 6, _safe(f"Most Common Schedule: {top_s} "
                                     f"({top_n:,} children, {top_p}%)"), ln=True)
            pdf.ln(5)

            # ── Chart sections ────────────────────────────────────────────────
            def _chart(section_title, key):
                path = charts.get(key)
                if not path or not os.path.exists(path):
                    return
                if pdf.get_y() > 200:
                    pdf.add_page()
                pdf.set_font("Helvetica", "B", 13)
                pdf.cell(0, 8, _safe(section_title), ln=True)
                pdf.ln(1)
                try:
                    pdf.image(path, x=15, w=180)
                except Exception:
                    pass
                pdf.ln(5)

            _chart("Immunization Schedule Distribution", "schedule_chart")
            _chart("Gender Distribution",                "gender_chart")
            _chart("Health Facilities Distribution",     "facility_chart")

            # ── Age distribution table ────────────────────────────────────────
            age = stats.get("age_distribution", {})
            if age:
                if pdf.get_y() > 210:
                    pdf.add_page()
                pdf.set_font("Helvetica", "B", 13)
                pdf.cell(0, 8, "Age Distribution", ln=True)
                pdf.ln(1)
                pdf.set_draw_color(*RED); pdf.set_line_width(0.4)
                w1, w2, w3 = 90, 45, 45
                pdf.set_font("Helvetica", "B", 10)
                pdf.cell(w1, 8, "Age Group",  border=1, align="C")
                pdf.cell(w2, 8, "Number",     border=1, align="C")
                pdf.cell(w3, 8, "Percentage", border=1, align="C", ln=True)
                pdf.set_font("Helvetica", "", 10)
                for band, val in age.items():
                    if isinstance(val, dict):
                        cnt, p = val.get("count", 0), val.get("percentage")
                    else:
                        cnt, p = val, None
                    if p is None:
                        p = round(cnt / total * 100, 1) if total else 0
                    pdf.cell(w1, 8, _safe(str(band)), border=1)
                    pdf.cell(w2, 8, f"{cnt:,}", border=1, align="C")
                    pdf.cell(w3, 8, f"{p}%",   border=1, align="C", ln=True)
                pdf.set_line_width(0.2); pdf.ln(6)

            # ── Required actions ──────────────────────────────────────────────
            if pdf.get_y() > 220:
                pdf.add_page()
            pdf.set_text_color(*RED); pdf.set_font("Helvetica", "B", 13)
            pdf.cell(0, 8, "REQUIRED ACTIONS", ln=True)
            pdf.set_text_color(0, 0, 0); pdf.ln(1)
            pdf.set_font("Helvetica", "", 10)
            for line in [
                "1. Provide immediate nutrition counseling for all identified stunted children",
                "2. Coordinate follow-up visits aligned with immunization schedules",
                "3. Engage community health workers for home visits and monitoring",
                "4. Send the list to the health facility to record them in the",
                "   e-Ubuzima system via the Growth monitoring module",
            ]:
                pdf.cell(0, 6, _safe(line), ln=True)
            pdf.cell(0, 6, "For technical support contact:", ln=True)
            pdf.cell(0, 6, _safe(f"{SUPPORT_CONTACT}: {SUPPORT_PHONE}"), ln=True)

            pdf.output(output_path)
            self.log.info(f"PDF saved: {output_path}")
            return True

        except Exception as exc:
            import traceback
            self.log.error(f"PDF error for {name}: {exc}\n{traceback.format_exc()}")
            return False

    # ── National overview (senior/ministry officials) ─────────────────────────

    def create_national_overview_pdf(
        self,
        stats: dict,
        charts: dict,
        year: int,
        month: int,
        output_path: str,
        total_vaccinated: int = 0,
        hospital_summary: list[dict] | None = None,
        total_assessed: int = 0,
        risk_summary: dict | None = None,
        district_summary: list[dict] | None = None,
    ) -> bool:
        """
        Create a NATIONAL-level overview PDF (all districts/hospitals combined) —
        for senior/ministry officials, distinct from the per-hospital/facility
        reports. Same visual language as create_pdf_report but with national
        totals, a trend chart, a province chart, a top-hospitals table, and an
        at-risk children summary instead of a single-facility breakdown.

        Section order is deliberately "story → action": trend first (where are
        we headed), geography (where's it worst), at-risk (who needs follow-up
        now), then priorities (what to do about it) — each section sets up the
        next rather than being an unordered dump of charts.

        risk_summary: output of core.risk_classifier.get_risk_summary(), or
        None to omit the At-Risk Children section entirely.
        """
        try:
            from config.app_config import SUPPORT_PHONE, SUPPORT_CONTACT

            period      = f"{MONTH_NAMES.get(month, str(month))} {year}"
            report_date = datetime.now().strftime("%B %d, %Y")
            total       = stats.get("total_stunted_children", 0)
            avg_age     = stats.get("avg_age_months")
            vacc        = total_vaccinated or 0
            # Rate is over ASSESSED children (usable height/DOB) — same as the
            # dashboard and the trend chart, so all three show the same %.
            denom       = total_assessed or vacc
            pct_vacc    = round(total / denom * 100, 1) if denom else 0
            RED         = (192, 57, 43)

            pdf = _PDF(title="National Stunting Surveillance Overview")
            pdf.alias_nb_pages()
            pdf.add_page()

            # ── Title ─────────────────────────────────────────────────────────
            pdf.set_text_color(*RED)
            pdf.set_font("Helvetica", "B", 18)
            pdf.cell(0, 12, _safe("NATIONAL STUNTING SURVEILLANCE OVERVIEW"),
                     ln=True, align="C")
            pdf.ln(1)
            pdf.set_text_color(0, 0, 0)
            pdf.set_font("Helvetica", "B", 13)
            pdf.cell(0, 8, _safe("All Districts — Rwanda"), ln=True, align="C")
            pdf.set_text_color(*RED)
            pdf.set_font("Helvetica", "", 10)
            pdf.cell(0, 6, _safe(f"Report Period: {period}"), ln=True, align="C")
            pdf.cell(0, 6, _safe(f"Report Date: {report_date}"), ln=True, align="C")
            pdf.set_text_color(0, 0, 0)
            pdf.ln(4)

            # ── Key statistics ────────────────────────────────────────────────
            bx, bw = 15, 180
            y0 = pdf.get_y()
            pdf.set_draw_color(*RED); pdf.set_line_width(0.6)
            pdf.set_fill_color(253, 237, 236)
            pdf.rect(bx, y0, bw, 26, "DF")
            pdf.set_xy(bx, y0 + 3)
            pdf.set_text_color(*RED); pdf.set_font("Helvetica", "B", 12)
            pdf.cell(bw, 7, "NATIONAL SUMMARY", ln=True, align="C")
            pdf.set_x(bx); pdf.set_font("Helvetica", "", 10)
            pdf.cell(bw, 6, _safe(f"Identified {total:,} stunted out of {denom:,} "
                                  f"children assessed ({pct_vacc}%) nationally"), ln=True, align="C")
            pdf.set_x(bx)
            n_hosp = len(hospital_summary or [])
            pdf.cell(bw, 6, _safe(f"Across {n_hosp} hospitals reporting this period"),
                     ln=True, align="C")
            pdf.set_text_color(0, 0, 0); pdf.set_line_width(0.2)
            pdf.set_y(y0 + 31)

            pdf.set_font("Helvetica", "B", 13)
            pdf.cell(0, 8, "Key Statistics", ln=True)
            pdf.ln(1)
            pdf.set_font("Helvetica", "", 11)
            pdf.cell(0, 6, _safe(f"Total Stunted Children (national): {total:,} out of "
                                 f"{denom:,} assessed ({pct_vacc}%)"), ln=True)
            if vacc:
                pdf.cell(0, 6, _safe(f"Children Vaccinated This Period: {vacc:,}"), ln=True)
            if avg_age:
                pdf.cell(0, 6, _safe(f"Average Age: {avg_age} months"), ln=True)
            pdf.ln(5)

            # ── Chart sections ────────────────────────────────────────────────
            def _chart(section_title, key):
                path = charts.get(key)
                if not path or not os.path.exists(path):
                    return
                if pdf.get_y() > 200:
                    pdf.add_page()
                pdf.set_font("Helvetica", "B", 13)
                pdf.cell(0, 8, _safe(section_title), ln=True)
                pdf.ln(1)
                try:
                    pdf.image(path, x=15, w=180)
                except Exception:
                    pass
                pdf.ln(5)

            _chart("Stunting Rate Over Time",        "trend_chart")
            _chart("Gender Distribution",           "gender_chart")
            _chart("Distribution by Province",       "province_chart")
            _chart("Top Hospitals by Stunted Cases", "top_hospitals_chart")

            # ── Age distribution table ────────────────────────────────────────
            age = stats.get("age_distribution", {})
            if age:
                if pdf.get_y() > 210:
                    pdf.add_page()
                pdf.set_font("Helvetica", "B", 13)
                pdf.cell(0, 8, "Age Distribution", ln=True)
                pdf.ln(1)
                pdf.set_draw_color(*RED); pdf.set_line_width(0.4)
                w1, w2, w3 = 90, 45, 45
                pdf.set_font("Helvetica", "B", 10)
                pdf.cell(w1, 8, "Age Group",  border=1, align="C")
                pdf.cell(w2, 8, "Number",     border=1, align="C")
                pdf.cell(w3, 8, "Percentage", border=1, align="C", ln=True)
                pdf.set_font("Helvetica", "", 10)
                for band, val in age.items():
                    if isinstance(val, dict):
                        cnt, p = val.get("count", 0), val.get("percentage")
                    else:
                        cnt, p = val, None
                    if p is None:
                        p = round(cnt / total * 100, 1) if total else 0
                    pdf.cell(w1, 8, _safe(str(band)), border=1)
                    pdf.cell(w2, 8, f"{cnt:,}", border=1, align="C")
                    pdf.cell(w3, 8, f"{p}%",   border=1, align="C", ln=True)
                pdf.set_line_width(0.2); pdf.ln(6)

            # ── Top hospitals table ────────────────────────────────────────────
            if hospital_summary:
                if pdf.get_y() > 190:
                    pdf.add_page()
                pdf.set_font("Helvetica", "B", 13)
                pdf.cell(0, 8, "Top 10 Hospitals by Stunted Children", ln=True)
                pdf.ln(1)
                pdf.set_draw_color(*RED); pdf.set_line_width(0.4)
                cw = [60, 40, 30, 30, 20]
                pdf.set_font("Helvetica", "B", 9)
                for w, h in zip(cw, ["Hospital", "District", "Total", "Severe", "Avg Age"]):
                    pdf.cell(w, 8, h, border=1, align="C")
                pdf.ln()
                pdf.set_font("Helvetica", "", 9)
                for row in hospital_summary[:10]:
                    pdf.cell(cw[0], 7, _safe(row.get("hospital", "")), border=1)
                    pdf.cell(cw[1], 7, _safe(row.get("district", "")), border=1)
                    pdf.cell(cw[2], 7, f"{row.get('total', 0):,}", border=1, align="C")
                    pdf.cell(cw[3], 7, f"{row.get('severe', 0):,}", border=1, align="C")
                    avg = row.get("avg_age")
                    pdf.cell(cw[4], 7, f"{avg}" if avg is not None else "—",
                            border=1, align="C", ln=True)
                pdf.set_line_width(0.2); pdf.ln(6)

            # ── District breakdown table (ALL districts, not just a top-N chart) ─
            if district_summary:
                if pdf.get_y() > 200:
                    pdf.add_page()
                pdf.set_font("Helvetica", "B", 13)
                pdf.cell(0, 8, _safe(f"Geographic Hotspots — All {len(district_summary)} "
                                     f"Districts (Stunting Prevalence)"), ln=True)
                pdf.ln(1)
                pdf.set_draw_color(*RED); pdf.set_line_width(0.4)
                cw = [65, 38, 38, 38]

                def _district_header():
                    pdf.set_font("Helvetica", "B", 9)
                    for w, h in zip(cw, ["District", "Assessed", "Stunted", "Prevalence %"]):
                        pdf.cell(w, 8, h, border=1, align="C")
                    pdf.ln()
                    pdf.set_font("Helvetica", "", 9)

                _district_header()
                for row in district_summary:
                    if pdf.get_y() > 275:
                        pdf.add_page()
                        _district_header()
                    pdf.cell(cw[0], 7, _safe(row.get("district", "")), border=1)
                    pdf.cell(cw[1], 7, f"{row.get('n', 0):,}", border=1, align="C")
                    pdf.cell(cw[2], 7, f"{row.get('stunted', 0):,}", border=1, align="C")
                    pdf.cell(cw[3], 7, f"{row.get('rate', 0)}%", border=1, align="C", ln=True)
                pdf.set_line_width(0.2); pdf.ln(6)

            # ── At-risk children ────────────────────────────────────────────────
            # Growth-velocity risk (distinct from stunting) — who needs follow-up
            # attention now, and briefly, how that was determined.
            if risk_summary:
                if pdf.get_y() > 210:
                    pdf.add_page()
                r_total = risk_summary.get("total", 0)
                r_high  = risk_summary.get("high", 0)
                r_med   = risk_summary.get("medium", 0)

                bx, bw = 15, 180
                y0 = pdf.get_y()
                pdf.set_draw_color(*RED); pdf.set_line_width(0.6)
                pdf.set_fill_color(255, 244, 230)
                pdf.rect(bx, y0, bw, 20, "DF")
                pdf.set_xy(bx, y0 + 3)
                pdf.set_text_color(*RED); pdf.set_font("Helvetica", "B", 12)
                pdf.cell(bw, 7, "AT-RISK CHILDREN", ln=True, align="C")
                pdf.set_x(bx); pdf.set_font("Helvetica", "", 10)
                pdf.cell(bw, 6, _safe(
                    f"{r_total:,} children flagged for growth-velocity follow-up "
                    f"({r_high:,} HIGH, {r_med:,} MEDIUM) across "
                    f"{risk_summary.get('hospitals', 0)} hospitals"),
                    ln=True, align="C")
                pdf.set_text_color(0, 0, 0); pdf.set_line_width(0.2)
                pdf.set_y(y0 + 25)

                pdf.set_font("Helvetica", "B", 10)
                pdf.cell(0, 6, "How this is measured (brief):", ln=True)
                pdf.set_font("Helvetica", "", 9)
                for line in [
                    "- HIGH: weight loss between consecutive visits, MUAC below 11.5cm, or severe wasting.",
                    "- MEDIUM: growth velocity below 50% of the WHO-expected minimum for the child's age",
                    "  (or height velocity below 30% of expected), or MUAC between 11.5-12.5cm.",
                    "- Velocity requires 2+ recorded visits; children with a single visit are screened",
                    "  by MUAC and wasting status only. Full case lists are in each hospital's own report.",
                ]:
                    pdf.cell(0, 5, _safe(line), ln=True)
                pdf.ln(4)

            # ── Required actions ──────────────────────────────────────────────
            if pdf.get_y() > 220:
                pdf.add_page()
            pdf.set_text_color(*RED); pdf.set_font("Helvetica", "B", 13)
            pdf.cell(0, 8, "NATIONAL PRIORITIES", ln=True)
            pdf.set_text_color(0, 0, 0); pdf.ln(1)
            pdf.set_font("Helvetica", "", 10)
            priorities = ["1. Prioritise nutrition interventions in the highest-burden districts/hospitals"]
            if risk_summary and risk_summary.get("total"):
                priorities.append(
                    f"2. Follow up the {risk_summary['total']:,} at-risk children above without delay "
                    f"({risk_summary.get('high', 0):,} of them HIGH risk)")
            else:
                priorities.append("2. Coordinate with district health officers on follow-up compliance")
            priorities.append(f"{len(priorities)+1}. Review facility-level reports for detailed case lists and contacts")
            for line in priorities:
                pdf.cell(0, 6, _safe(line), ln=True)
            pdf.cell(0, 6, "For technical support contact:", ln=True)
            pdf.cell(0, 6, _safe(f"{SUPPORT_CONTACT}: {SUPPORT_PHONE}"), ln=True)

            pdf.output(output_path)
            self.log.info(f"National overview PDF saved: {output_path}")
            return True

        except Exception as exc:
            import traceback
            self.log.error(f"National overview PDF error: {exc}\n{traceback.format_exc()}")
            return False

    # ── Excel case list ────────────────────────────────────────────────────────

    def create_excel_report(
        self,
        df,
        output_path: str,
        sheet_name: str = "Stunted Children",
    ) -> bool:
        """
        Create an Excel case list from a stunted children DataFrame.
        Returns True on success.
        """
        import pandas as pd
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = sheet_name

            # Resolve columns present in df
            out_cols = [(src, hdr) for src, hdr in CASE_LIST_COLS if src in df.columns]
            if not out_cols:
                out_cols = [(c, c) for c in df.columns[:20]]

            # Style helpers
            hdr_font  = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
            hdr_fill  = PatternFill("solid", fgColor=_DARK_HEX)
            thin_side = Side(style="thin", color="D5D8DC")
            thin_bdr  = Border(left=thin_side, right=thin_side,
                               top=thin_side, bottom=thin_side)
            alt_fill  = PatternFill("solid", fgColor="F2F3F4")
            center    = Alignment(horizontal="center", vertical="center")
            wrap      = Alignment(wrap_text=True, vertical="top")

            # Header row
            for col_idx, (_, hdr) in enumerate(out_cols, start=1):
                cell = ws.cell(row=1, column=col_idx, value=hdr)
                cell.font   = hdr_font
                cell.fill   = hdr_fill
                cell.border = thin_bdr
                cell.alignment = center

            ws.row_dimensions[1].height = 20

            # Data rows
            df_out = df[[s for s, _ in out_cols]].copy()
            df_out.columns = [h for _, h in out_cols]

            for row_idx, row in enumerate(df_out.itertuples(index=False), start=2):
                fill = alt_fill if row_idx % 2 == 0 else None
                for col_idx, val in enumerate(row, start=1):
                    # Convert numpy types
                    if hasattr(val, "item"):
                        val = val.item()
                    if pd.isna(val) if hasattr(pd, "isna") else val != val:
                        val = ""
                    cell = ws.cell(row=row_idx, column=col_idx, value=val)
                    cell.border = thin_bdr
                    cell.alignment = wrap
                    if fill:
                        cell.fill = fill

            # Auto-width (cap at 40)
            for col_idx in range(1, len(out_cols) + 1):
                col_letter = get_column_letter(col_idx)
                max_len = max(
                    (
                        len(str(ws.cell(r, col_idx).value or ""))
                        for r in range(1, min(ws.max_row + 1, 500))
                    ),
                    default=10,
                )
                ws.column_dimensions[col_letter].width = min(max_len + 2, 40)

            # Freeze header
            ws.freeze_panes = "A2"

            wb.save(output_path)
            self.log.info(f"Excel saved: {output_path} ({len(df_out):,} rows)")
            return True

        except Exception as exc:
            self.log.error(f"Excel error: {exc}")
            return False

    # ── At-risk children list (hospital report attachment) ─────────────────────

    AT_RISK_COLS = [
        ("entity_id",           "Entity ID"),
        ("child_name",          "Child Name"),
        ("gender",               "Gender"),
        ("age_in_months",       "Age (Months)"),
        ("risk_level",          "Risk Level"),
        ("weight_velocity",     "Weight Velocity (kg/mo)"),
        ("height_velocity",     "Height Velocity (cm/mo)"),
        ("muac_cm",             "MUAC (cm)"),
        ("wasting_status",      "Wasting Status"),
        ("health_facility",     "Health Facility"),
        ("mother_names",        "Mother Name"),
        ("mother_phone",        "Mother Phone"),
        ("father_names",        "Father Name"),
        ("father_phone",        "Father Phone"),
    ]

    def create_at_risk_excel(self, at_risk_df, output_path: str) -> bool:
        """
        Excel list of at-risk children (HIGH/MEDIUM growth-velocity risk) for
        one hospital — attached alongside the usual stunting report/case list,
        NOT a replacement for it. See core/risk_classifier.py's
        classify_at_risk() for how risk_level is determined.
        """
        import pandas as pd
        if at_risk_df is None or at_risk_df.empty:
            return False
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "At-Risk Children"

            out_cols = [(src, hdr) for src, hdr in self.AT_RISK_COLS if src in at_risk_df.columns]
            if not out_cols:
                return False

            hdr_font  = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
            hdr_fill  = PatternFill("solid", fgColor=_RED_HEX)
            thin_side = Side(style="thin", color="D5D8DC")
            thin_bdr  = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
            high_fill = PatternFill("solid", fgColor="FDEDEC")
            med_fill  = PatternFill("solid", fgColor="FEF9E7")
            center    = Alignment(horizontal="center", vertical="center")
            wrap      = Alignment(wrap_text=True, vertical="top")

            for col_idx, (_, hdr) in enumerate(out_cols, start=1):
                cell = ws.cell(row=1, column=col_idx, value=hdr)
                cell.font, cell.fill, cell.border, cell.alignment = hdr_font, hdr_fill, thin_bdr, center
            ws.row_dimensions[1].height = 20

            # Sort HIGH first so the most urgent cases are at the top.
            df_sorted = at_risk_df.copy()
            if "risk_level" in df_sorted.columns:
                df_sorted["_sort"] = df_sorted["risk_level"].map({"HIGH": 0, "MEDIUM": 1}).fillna(2)
                df_sorted = df_sorted.sort_values("_sort").drop(columns="_sort")

            df_out = df_sorted[[s for s, _ in out_cols]].copy()
            df_out.columns = [h for _, h in out_cols]
            risk_col_idx = next((i for i, (s, _) in enumerate(out_cols, start=1) if s == "risk_level"), None)

            for row_idx, row in enumerate(df_out.itertuples(index=False), start=2):
                risk_val = row[risk_col_idx - 1] if risk_col_idx else None
                fill = high_fill if risk_val == "HIGH" else med_fill if risk_val == "MEDIUM" else None
                for col_idx, val in enumerate(row, start=1):
                    if hasattr(val, "item"):
                        val = val.item()
                    if pd.isna(val) if hasattr(pd, "isna") else val != val:
                        val = ""
                    cell = ws.cell(row=row_idx, column=col_idx, value=val)
                    cell.border = thin_bdr
                    cell.alignment = wrap
                    if fill:
                        cell.fill = fill

            for col_idx in range(1, len(out_cols) + 1):
                col_letter = get_column_letter(col_idx)
                max_len = max(
                    (len(str(ws.cell(r, col_idx).value or ""))
                     for r in range(1, min(ws.max_row + 1, 500))),
                    default=10,
                )
                ws.column_dimensions[col_letter].width = min(max_len + 2, 35)

            ws.freeze_panes = "A2"
            wb.save(output_path)
            self.log.info(f"At-risk Excel saved: {output_path} ({len(df_out):,} children)")
            return True

        except Exception as exc:
            self.log.error(f"At-risk Excel error: {exc}")
            return False
